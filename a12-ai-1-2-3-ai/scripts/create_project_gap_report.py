from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "TeachNova工程缺陷与智能体差距分析.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_run_font(run, name="Heiti SC", size=None, bold=None, color=None):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    for key in ("w:eastAsia", "w:ascii", "w:hAnsi"):
        r_fonts.set(qn(key), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(document, text, level=1):
    p = document.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, name="Heiti SC", color="000000", bold=True)
    return p


def add_paragraph(document, text, bold_lead=None):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True, size=11)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11)
    return p


def add_bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(item)
        set_run_font(run, size=10.5)


def add_matrix(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_shading(hdr[idx], "1F4E79")
        set_cell_margins(hdr[idx])
        run = hdr[idx].paragraphs[0].add_run(header)
        set_run_font(run, size=10.5, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_margins(cells[idx])
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[idx], "F6FAFF")
            run = cells[idx].paragraphs[0].add_run(str(value))
            set_run_font(run, size=10)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    document.add_paragraph()
    return table


def build_doc():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = styles[style_name]
        style.font.name = "Heiti SC"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Heiti SC")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(11)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    title_run = title.add_run("TeachNova 工程缺陷与智能体差距分析")
    set_run_font(title_run, size=20, bold=True, color="000000")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("基于当前工程实现状态与最初赛题要求的对照评估")
    set_run_font(subtitle_run, size=11, color="4B5563")

    add_paragraph(
        document,
        "结论：当前工程已经跑通 AI 教学内容生成平台的主流程，包括资料上传、文档解析、RAG 检索、模型配置、聊天、大纲、Word、PPT 与反馈再生成。但从赛题定义看，它仍更接近教学内容生成系统原型，还没有完全达到多模态互动式教学智能体的水平。",
        "结论：",
    )

    add_heading(document, "一 当前已具备的工程能力")
    add_matrix(
        document,
        ["模块", "当前能力", "成熟度"],
        [
            ["前端工作台", "React Vite Ant Design，已形成 TeachNova 品牌、浅色玻璃风格和多页面工作台", "中"],
            ["后端服务", "FastAPI 提供任务、上传、解析、知识库、聊天、生成、迭代等接口", "中"],
            ["模型配置", "支持 DeepSeek Qwen Xiaomi MiMo 的保存、切换和测试连接", "中"],
            ["资料解析", "支持 PDF Word PPT，优先 MarkItDown，失败后走本地解析器", "中"],
            ["RAG 知识库", "Markdown 切片，Chroma 向量库，BGE 本地 Embedding，DashScope 预留，hash 兜底", "中"],
            ["课件生成", "可生成结构化大纲、Word 教案、PPT 课件，并输出质量检查报告", "中"],
            ["反馈迭代", "可输入教师反馈并重新生成大纲、Word、PPT 版本产物", "中"],
        ],
        widths=[1.45, 4.6, 0.9],
    )

    add_heading(document, "二 主要功能缺陷")
    add_bullets(
        document,
        [
            "本地知识库内容太少。当前正式入库主要是 TCP 三次握手测试切片，不能支撑复杂课程设计。",
            "主动追问能力不足。系统能抽取教学意图，但缺少明确的缺失字段判断和按字段追问策略。",
            "多模态能力不完整。PDF Word PPT 已支持，但图片 OCR、视频关键帧、音频转写和语音输入还没有形成闭环。",
            "PPT 生成质量偏基础。当前可以生成可编辑 PPT，但版式、图示、页面类型和视觉丰富度还不够强。",
            "Word 教案偏骨架化。缺少学情分析、教学资源、评价设计、板书设计、分层作业和素养目标等更专业内容。",
            "反馈迭代粒度较粗。已经支持再生成，但还不能稳定做到只修改某一页、某个案例或某段讲稿。",
            "工具调用日志不完整。前端已有可解释工作流，但还没有完整展示模型、解析器、向量库、生成器和质量检查工具的真实调用记录。",
            "模型失败降级不够清晰。连接失败、超时、API 格式差异和不同模型能力边界需要更明确的状态提示和兜底策略。",
        ],
    )

    add_heading(document, "三 与真正智能体的差距")
    add_matrix(
        document,
        ["能力维度", "当前工程状态", "真正智能体应具备的表现"],
        [
            ["目标驱动", "用户需要手动切换页面和点击生成按钮", "能根据教师目标自动判断下一步动作"],
            ["自主规划", "流程是上传 解析 检索 生成的接口串联", "能生成并执行任务计划，必要时暂停追问"],
            ["工具调度", "后端有多个服务，但缺统一调度层", "能动态选择解析器、检索器、生成器和质检工具"],
            ["主动追问", "依赖 prompt 和简单 fallback", "按缺失字段精准追问并确认需求完整度"],
            ["状态记忆", "任务字段保存了一部分状态", "能清楚知道当前任务阶段、历史版本、反馈修改点"],
            ["可验证行动", "引用和 thinking 摘要已有雏形", "能展示每一步调用了什么工具、用了哪些资料、改了哪些产物"],
            ["多模态理解", "主要是文档解析", "能理解图片、视频、音频，并把内容用于生成"],
        ],
        widths=[1.4, 2.75, 2.75],
    )

    add_heading(document, "四 赛题要求对照")
    add_matrix(
        document,
        ["赛题要求", "当前达成情况", "风险"],
        [
            ["本地知识库 RAG", "流程已通，BGE 与 Chroma 已接入", "资料量不足，检索质量难体现"],
            ["语音和文字输入", "文字输入已实现", "语音输入未实现"],
            ["主动澄清需求", "部分实现", "缺少稳定追问状态机"],
            ["PDF Word PPT 图片 视频资料融合", "PDF Word PPT 已实现", "图片和视频未完成"],
            ["生成 PPT 和 Word 教案", "已实现", "质量和模板专业度仍需提升"],
            ["动画创意或互动小游戏", "可在大纲中生成互动描述", "还没有真正 HTML5 小游戏或动画导出"],
            ["反馈再生成", "已实现版本化再生成", "局部编辑能力不足"],
        ],
        widths=[2.05, 3.0, 1.85],
    )

    add_heading(document, "五 优先改进路线")
    add_bullets(
        document,
        [
            "第一优先级：新增 Agent Orchestrator。把意图判断、追问、RAG、生成、迭代统一到一个智能体调度层。",
            "第二优先级：强化主动追问状态机。围绕课程主题、授课对象、课时、知识点、重点难点、互动形式、产出类型做缺失字段检测。",
            "第三优先级：补充正式本地知识库。至少准备 20 份高质量课程资料，覆盖计算机网络、Python、数据库和 AI 基础。",
            "第四优先级：补图片 OCR 或视频摘要。建议先做图片 OCR，成本低，能快速补齐多模态短板。",
            "第五优先级：升级 PPT 生成。为大纲增加 slide_type，并按封面、目录、概念、流程、案例、互动、总结分别设计版式。",
            "第六优先级：完善工具调用日志。把模型、解析器、检索 collection、命中 chunk、生成文件和质检结果展示给用户。",
        ],
    )

    add_heading(document, "六 建议的下一步工程任务")
    add_paragraph(
        document,
        "建议下一步先实现 Agent Orchestrator 与主动追问状态机。这一改动对界面影响较小，但会显著提升系统的智能体感，使它从多个功能模块的集合变成一个围绕教师目标自主推进的教学智能体。",
    )
    add_bullets(
        document,
        [
            "新增 backend/app/services/agent_orchestrator.py。",
            "定义任务阶段：clarify retrieve generate review iterate。",
            "定义缺失字段：课程主题、授课对象、课时、知识点、重点难点、互动形式、产物类型。",
            "聊天接口先调用 Orchestrator，再决定追问、检索、生成或等待反馈。",
            "前端展示真实工作流步骤和当前任务阶段。",
        ],
    )

    document.core_properties.title = "TeachNova 工程缺陷与智能体差距分析"
    document.core_properties.subject = "项目评估"
    document.core_properties.author = "TeachNova 项目组"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()
