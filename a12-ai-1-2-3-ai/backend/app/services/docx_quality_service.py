import json
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document

from app.core.config import settings


REQUIRED_DOCX_SECTIONS = [
    "基本信息",
    "一、教学目标",
    "二、教学重点与难点",
    "三、教学过程",
    "四、课件页面设计",
    "五、课堂活动设计",
    "六、课后作业",
]


@dataclass
class DocxQualityReport:
    status: str
    passed: bool
    docx_path: str
    report_path: str | None
    checks: dict[str, bool]
    warnings: list[str]
    text_preview: str
    paragraph_count: int
    table_count: int


def inspect_docx_quality(docx_path: Path) -> DocxQualityReport:
    warnings: list[str] = []
    checks = {
        "zip_integrity": _check_zip_integrity(docx_path, warnings),
        "readable_text": False,
        "required_sections": False,
        "has_tables": False,
        "has_chinese_font_markers": _check_chinese_font_markers(docx_path, warnings),
    }

    text = ""
    paragraph_count = 0
    table_count = 0
    try:
        document = Document(docx_path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text = _extract_table_text(document)
        text = "\n".join([*paragraphs, *table_text]).strip()
        paragraph_count = len(paragraphs)
        table_count = len(document.tables)
        checks["readable_text"] = len(text) >= 20
        checks["required_sections"] = all(section in text for section in REQUIRED_DOCX_SECTIONS)
        checks["has_tables"] = table_count > 0
    except Exception as exc:
        warnings.append(f"DOCX 内容读取失败：{exc}")

    for check_name, ok in checks.items():
        if not ok:
            warnings.append(f"质量检查未通过：{check_name}")

    passed = all(checks.values())
    return DocxQualityReport(
        status="passed" if passed else "warning",
        passed=passed,
        docx_path=str(docx_path),
        report_path=None,
        checks=checks,
        warnings=warnings,
        text_preview=text[:400],
        paragraph_count=paragraph_count,
        table_count=table_count,
    )


def save_docx_quality_report(task_id: int, report: DocxQualityReport) -> Path:
    report_dir = settings.outputs_dir / "docx" / "qa"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / f"lesson_plan_task_{task_id}_qa.json"
    report.report_path = str(report_path)
    with report_path.open("w", encoding="utf-8") as file:
        json.dump(asdict(report), file, ensure_ascii=False, indent=2)
    return report_path


def docx_quality_to_dict(report: DocxQualityReport) -> dict[str, Any]:
    return asdict(report)


def _check_zip_integrity(docx_path: Path, warnings: list[str]) -> bool:
    try:
        with zipfile.ZipFile(docx_path) as archive:
            bad_file = archive.testzip()
        if bad_file:
            warnings.append(f"DOCX 压缩包存在损坏文件：{bad_file}")
            return False
        return True
    except Exception as exc:
        warnings.append(f"DOCX 压缩包校验失败：{exc}")
        return False


def _check_chinese_font_markers(docx_path: Path, warnings: list[str]) -> bool:
    try:
        with zipfile.ZipFile(docx_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
            styles_xml = archive.read("word/styles.xml").decode("utf-8", errors="ignore")
    except Exception as exc:
        warnings.append(f"DOCX 字体标记读取失败：{exc}")
        return False

    xml = f"{document_xml}\n{styles_xml}"
    return "w:eastAsia" in xml and "zh-CN" in xml


def _extract_table_text(document: Document) -> list[str]:
    values: list[str] = []
    for table in document.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                values.append(row_text)
    return values
