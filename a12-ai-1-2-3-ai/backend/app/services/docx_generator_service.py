from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from app.core.config import settings
from app.models.task import CourseTask
from app.services.lesson_plan_service import load_or_create_lesson_plan
from app.utils.data import as_object_list, as_string_list


DOCX_FONT = "Songti SC"


def generate_docx_lesson_plan(task: CourseTask) -> tuple[Path, dict[str, Any]]:
    plan = load_or_create_lesson_plan(task)
    output_path = settings.outputs_dir / "docx" / f"lesson_plan_task_{task.id}.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _setup_document(document)
    _add_title(document, plan)
    _add_basic_info(document, plan, task)
    _add_list_section(document, "一、教学目标", plan.get("teaching_targets", []))
    _add_list_section(document, "二、教学重点与难点", plan.get("key_difficulties", []))
    _add_process_section(document, plan.get("teaching_process", []))
    _add_slides_section(document, plan.get("slides", []))
    _add_activities_section(document, plan.get("activities", []))
    _add_list_section(document, "六、课后作业", plan.get("homework", []))
    _add_references_section(document, plan.get("references", []))
    _apply_document_fonts(document)

    document.save(output_path)
    return output_path, plan


def _setup_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Pt(56)
    section.bottom_margin = Pt(56)
    section.left_margin = Pt(64)
    section.right_margin = Pt(64)

    styles = document.styles
    styles["Normal"].font.name = DOCX_FONT
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        styles[style_name].font.name = DOCX_FONT
        _set_style_font(styles[style_name], DOCX_FONT)
    _set_style_font(styles["Normal"], DOCX_FONT)


def _set_style_font(style: Any, font_name: str) -> None:
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font_name)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font_name)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hint"), "eastAsia")


def _apply_document_fonts(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            _set_run_font(run, DOCX_FONT)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_font(run, DOCX_FONT)


def _set_run_font(run: Any, font_name: str) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _add_title(document: Document, plan: dict[str, Any]) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(f"{plan.get('title', '课程')} 教案")

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("由 TeachNova 根据结构化课件大纲自动生成").italic = True


def _add_basic_info(document: Document, plan: dict[str, Any], task: CourseTask) -> None:
    document.add_heading("基本信息", level=1)
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    rows = [
        ("课程主题", str(plan.get("title") or task.title)),
        ("授课对象", str(plan.get("audience") or task.audience or "待确认")),
        ("课时长度", f"{plan.get('duration_minutes') or task.duration_minutes or '待确认'} 分钟"),
        ("课程状态", task.intent_status or "draft"),
    ]
    for row, (label, value) in zip(table.rows, rows, strict=False):
        row.cells[0].text = label
        row.cells[1].text = value


def _add_list_section(document: Document, title: str, items: Any) -> None:
    document.add_heading(title, level=1)
    values = as_string_list(items)
    if not values:
        document.add_paragraph("待补充。")
        return
    for item in values:
        document.add_paragraph(item, style="List Bullet")


def _add_process_section(document: Document, process: Any) -> None:
    document.add_heading("三、教学过程", level=1)
    rows = as_object_list(process)
    if not rows:
        document.add_paragraph("待补充。")
        return

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["环节", "时间", "教师活动", "学生活动"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for item in rows:
        row = table.add_row().cells
        row[0].text = str(item.get("phase") or "")
        minutes = item.get("minutes")
        row[1].text = f"{minutes} 分钟" if minutes else ""
        row[2].text = str(item.get("teacher_activity") or "")
        row[3].text = str(item.get("student_activity") or "")


def _add_slides_section(document: Document, slides: Any) -> None:
    document.add_heading("四、课件页面设计", level=1)
    slide_items = as_object_list(slides)
    if not slide_items:
        document.add_paragraph("待补充。")
        return

    for index, slide in enumerate(slide_items, start=1):
        document.add_heading(f"第 {index} 页：{slide.get('title') or '内容页'}", level=2)
        for point in as_string_list(slide.get("bullet_points")):
            document.add_paragraph(point, style="List Bullet")
        if slide.get("speaker_notes"):
            document.add_paragraph(f"讲授提示：{slide['speaker_notes']}")
        if slide.get("visual_suggestion"):
            document.add_paragraph(f"视觉建议：{slide['visual_suggestion']}")
        if slide.get("interaction"):
            document.add_paragraph(f"互动设计：{slide['interaction']}")


def _add_activities_section(document: Document, activities: Any) -> None:
    document.add_heading("五、课堂活动设计", level=1)
    rows = as_object_list(activities)
    if not rows:
        document.add_paragraph("待补充。")
        return

    for item in rows:
        document.add_heading(str(item.get("name") or "课堂活动"), level=2)
        document.add_paragraph(f"活动类型：{item.get('type') or '待确认'}")
        document.add_paragraph(f"活动说明：{item.get('description') or '待补充。'}")
        document.add_paragraph(f"预期效果：{item.get('expected_outcome') or '待补充。'}")


def _add_references_section(document: Document, references: Any) -> None:
    document.add_heading("七、参考资料", level=1)
    rows = as_object_list(references)
    if not rows:
        document.add_paragraph("暂无引用资料。")
        return
    for item in rows:
        source = item.get("source") or "本地知识库"
        page = f"，页码 {item['page']}" if item.get("page") else ""
        usage = item.get("usage") or "生成课件内容参考"
        document.add_paragraph(f"{source}{page}：{usage}", style="List Bullet")
