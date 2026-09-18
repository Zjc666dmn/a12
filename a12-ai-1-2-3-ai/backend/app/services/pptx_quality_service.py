import json
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from pptx import Presentation

from app.core.config import settings


REQUIRED_PPTX_KEYWORDS = ["学习目标", "重点", "总结"]


@dataclass
class PptxQualityReport:
    status: str
    passed: bool
    pptx_path: str
    report_path: str | None
    checks: dict[str, bool]
    warnings: list[str]
    slide_count: int
    text_preview: str


def inspect_pptx_quality(pptx_path: Path, expected_title: str | None = None) -> PptxQualityReport:
    warnings: list[str] = []
    checks = {
        "zip_integrity": _check_zip_integrity(pptx_path, warnings),
        "has_slides": False,
        "has_title": False,
        "has_required_sections": False,
        "no_empty_slides": False,
        "has_speaker_notes": _check_speaker_notes(pptx_path, warnings),
    }

    slide_count = 0
    text = ""
    try:
        presentation = Presentation(pptx_path)
        slide_count = len(presentation.slides)
        slide_texts = [_extract_slide_text(slide) for slide in presentation.slides]
        text = "\n".join(item for item in slide_texts if item).strip()
        checks["has_slides"] = slide_count >= 5
        checks["has_title"] = bool(expected_title and expected_title in text) or bool(text)
        checks["has_required_sections"] = all(keyword in text for keyword in REQUIRED_PPTX_KEYWORDS)
        checks["no_empty_slides"] = all(bool(item.strip()) for item in slide_texts)
    except Exception as exc:
        warnings.append(f"PPTX 内容读取失败：{exc}")

    for check_name, ok in checks.items():
        if not ok:
            warnings.append(f"质量检查未通过：{check_name}")

    passed = all(checks.values())
    return PptxQualityReport(
        status="passed" if passed else "warning",
        passed=passed,
        pptx_path=str(pptx_path),
        report_path=None,
        checks=checks,
        warnings=warnings,
        slide_count=slide_count,
        text_preview=text[:500],
    )


def save_pptx_quality_report(task_id: int, report: PptxQualityReport) -> Path:
    report_dir = settings.outputs_dir / "ppt" / "qa"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / f"lesson_slides_task_{task_id}_qa.json"
    report.report_path = str(report_path)
    with report_path.open("w", encoding="utf-8") as file:
        json.dump(asdict(report), file, ensure_ascii=False, indent=2)
    return report_path


def pptx_quality_to_dict(report: PptxQualityReport) -> dict[str, Any]:
    return asdict(report)


def _check_zip_integrity(pptx_path: Path, warnings: list[str]) -> bool:
    try:
        with zipfile.ZipFile(pptx_path) as archive:
            bad_file = archive.testzip()
        if bad_file:
            warnings.append(f"PPTX 压缩包存在损坏文件：{bad_file}")
            return False
        return True
    except Exception as exc:
        warnings.append(f"PPTX 压缩包校验失败：{exc}")
        return False


def _check_speaker_notes(pptx_path: Path, warnings: list[str]) -> bool:
    try:
        with zipfile.ZipFile(pptx_path) as archive:
            note_files = [name for name in archive.namelist() if name.startswith("ppt/notesSlides/notesSlide")]
    except Exception as exc:
        warnings.append(f"PPTX 备注读取失败：{exc}")
        return False
    return bool(note_files)


def _extract_slide_text(slide: Any) -> str:
    values: list[str] = []
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            text = str(shape.text).strip()
            if text:
                values.append(text)
    return "\n".join(values)
