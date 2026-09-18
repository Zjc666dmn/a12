from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from pptx import Presentation

from app.services.markitdown_parser import MarkItDownParseError, parse_with_markitdown


def parse_document(file_path: str | Path, file_type: str | None = None) -> dict[str, Any]:
    path = Path(file_path)
    suffix = path.suffix.lower()

    try:
        return parse_with_markitdown(path, file_type)
    except MarkItDownParseError:
        pass

    if suffix == ".pdf" or file_type == "application/pdf":
        return parse_pdf(path)

    if suffix == ".docx" or file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return parse_docx(path)

    if suffix == ".pptx" or file_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        return parse_pptx(path)

    if suffix == ".doc":
        raise ValueError("暂不支持解析 .doc 旧版 Word 文件，请转换为 .docx 后上传")

    raise ValueError("仅支持解析 PDF、Word（.docx）和 PPT（.pptx）文件")


def parse_pdf(file_path: str | Path) -> dict[str, Any]:
    path = Path(file_path)
    pages: list[dict[str, Any]] = []

    with pymupdf.open(path) as document:
        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            pages.append({"page": index, "text": text})

    full_text = "\n\n".join(page["text"] for page in pages if page["text"])
    return {
        "title": path.stem,
        "file_type": "pdf",
        "pages": pages,
        "text": full_text,
        "metadata": {
            "page_count": len(pages),
            "character_count": len(full_text),
        },
    }


def parse_docx(file_path: str | Path) -> dict[str, Any]:
    path = Path(file_path)
    document = Document(path)
    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    full_text = "\n".join(blocks)
    return {
        "title": path.stem,
        "file_type": "docx",
        "pages": [{"page": 1, "text": full_text}],
        "text": full_text,
        "metadata": {
            "paragraph_count": len([paragraph for paragraph in document.paragraphs if paragraph.text.strip()]),
            "table_count": len(document.tables),
            "character_count": len(full_text),
        },
    }


def parse_pptx(file_path: str | Path) -> dict[str, Any]:
    path = Path(file_path)
    presentation = Presentation(path)
    pages: list[dict[str, Any]] = []

    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        tables: list[dict[str, Any]] = []
        title = ""

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    if not title:
                        title = text.splitlines()[0]
                    texts.append(text)

            if getattr(shape, "has_table", False):
                rows: list[list[str]] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        rows.append(cells)
                if rows:
                    table_text = "\n".join(" | ".join(row) for row in rows)
                    tables.append({"rows": rows, "text": table_text})
                    texts.append(table_text)

        slide_text = "\n".join(texts)
        pages.append(
            {
                "page": index,
                "title": title,
                "text": slide_text,
                "tables": tables,
                "images": [],
            }
        )

    full_text = "\n\n".join(page["text"] for page in pages if page["text"])
    return {
        "title": path.stem,
        "file_type": "pptx",
        "pages": pages,
        "text": full_text,
        "metadata": {
            "slide_count": len(pages),
            "character_count": len(full_text),
        },
    }
