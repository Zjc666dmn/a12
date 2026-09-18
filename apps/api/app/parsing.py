from __future__ import annotations

import io
import os
from pathlib import Path
from typing import IO

from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.models import Chunk, Document, DocumentPage, DocumentSection
from app.search_indexing import token_count


def _extension(path: str) -> str:
    return Path(path).suffix.lower()


def _read_pdf(path: str) -> tuple[str, list[str]]:
    import fitz

    pdf = fitz.open(path)
    full_text_parts: list[str] = []
    pages: list[str] = []
    for page in pdf:
        text = page.get_text("text")
        pages.append(text)
        full_text_parts.append(text)
    return "\n\n".join(full_text_parts), pages


def _read_docx(path: str) -> tuple[str, list[str]]:
    from docx import Document as Docx

    doc = Docx(path)
    full_text = "\n\n".join(p.text for p in doc.paragraphs)
    pages = [full_text]
    return full_text, pages


def _read_pptx(path: str) -> tuple[str, list[str]]:
    from pptx import Presentation

    prs = Presentation(path)
    slides: list[str] = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                parts.append(shape.text)
        slides.append("\n".join(parts))
    return "\n\n".join(slides), slides


def _read_plain(path: str) -> tuple[str, list[str]]:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    return text, [text]


def read_document(path: str) -> tuple[str, list[str], str]:
    ext = _extension(path)
    if ext == ".pdf":
        text, pages = _read_pdf(path)
        return text, pages, "pdf"
    if ext == ".docx":
        text, pages = _read_docx(path)
        return text, pages, "docx"
    if ext == ".pptx":
        text, pages = _read_pptx(path)
        return text, pages, "pptx"
    text, pages = _read_plain(path)
    return text, pages, "txt"


def chunk_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []
    start = 0
    chunks: list[str] = []
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks


def create_initial_page_texts(text: str, pages: list[str]) -> list[str]:
    if len(pages) <= 1:
        return [text]
    return pages


async def parse_document(session: AsyncSession, document: Document) -> None:
    document.status = "parsing"
    document.status_message = "Parsing started"
    await session.flush()

    full_text, pages, file_type = read_document(document.file_name)
    if not full_text.strip():
        document.status = "failed"
        document.status_message = "Document content is empty"
        await session.flush()
        return

    page_texts = create_initial_page_texts(full_text, pages)
    for idx, page_text in enumerate(page_texts, start=1):
        session.add(
            DocumentPage(
                document_id=document.id,
                page_number=idx,
                text=page_text[:100_000],
                has_table="table" in page_text.lower(),
                has_formula="=" in page_text or "$" in page_text,
            )
        )

    raw_text_dir = Path(settings.upload_dir) / str(document.id)
    raw_text_dir.mkdir(parents=True, exist_ok=True)
    (raw_text_dir / "full_text.txt").write_text(full_text, encoding="utf-8")

    sections: list[DocumentSection] = []
    for heading in [p for p in page_texts[:5] if p.strip()][:3]:
        section = DocumentSection(
            document_id=document.id,
            level=1,
            order_index=len(sections) + 1,
            title=heading.split("\n", 1)[0][:260] or f"Section {len(sections) + 1}",
            content_hash=str(hash(heading)),
        )
        session.add(section)
        sections.append(section)

    chunks: list[Chunk] = []
    for idx, chunk_text_value in enumerate(chunk_text(full_text, settings.chunk_size, settings.chunk_overlap)):
        chunks.append(
            Chunk(
                document_id=document.id,
                page_start=1,
                page_end=max(1, len(pages)),
                chunk_index=idx,
                content=chunk_text_value,
                content_hash=str(hash(chunk_text_value)),
                token_count=token_count(chunk_text_value),
                stage=document.stage,
                grade=document.grade,
                semester=document.semester,
                subject=document.subject,
                publisher=document.publisher,
                textbook_version=document.textbook_version,
                book_name=document.book_name,
                chapter=document.chapter,
                section=document.section,
                resource_type=document.resource_type,
            )
        )
    session.add_all(chunks)
    document.status = "completed"
    document.status_message = f"Parsed {len(chunks)} chunks"
    await session.flush()
